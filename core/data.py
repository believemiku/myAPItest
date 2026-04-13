import os
import csv
import json
import yaml
from openpyxl import load_workbook
from typing import List, Dict, Any, Optional

class DataManager:
    def __init__(self):
        self.data_cache = {}  # 数据文件缓存，键为文件路径，值为加载的数据
    
    def load_data(self, data_path: str) -> List[Dict[str, Any]]:
        """加载测试数据"""
        # 转换为绝对路径，确保缓存键的唯一性
        data_path = os.path.abspath(data_path)
        
        # 检查缓存
        if data_path in self.data_cache:
            return self.data_cache[data_path]
        
        if not os.path.exists(data_path):
            raise Exception(f"数据文件不存在: {data_path}")
        
        file_ext = os.path.splitext(data_path)[1].lower()
        
        if file_ext == '.xlsx' or file_ext == '.xls':
            data = self._load_excel_data(data_path)
        elif file_ext == '.csv':
            data = self._load_csv_data(data_path)
        elif file_ext == '.json':
            data = self._load_json_data(data_path)
        elif file_ext == '.yaml' or file_ext == '.yml':
            data = self._load_yaml_data(data_path)
        else:
            raise Exception(f"不支持的数据文件格式: {file_ext}")
        
        # 缓存数据
        self.data_cache[data_path] = data
        return data
    
    def _load_excel_data(self, file_path: str) -> List[Dict[str, Any]]:
        """加载Excel数据"""
        workbook = load_workbook(file_path, read_only=True)  # 使用只读模式提高性能
        sheet = workbook.active
        
        # 读取表头
        headers = []
        for cell in sheet[1]:
            headers.append(cell.value)
        
        # 读取数据
        data_list = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_data = {}
            for i, value in enumerate(row):
                if i < len(headers):
                    row_data[headers[i]] = value
            data_list.append(row_data)
        
        return data_list
    
    def _load_csv_data(self, file_path: str) -> List[Dict[str, Any]]:
        """加载CSV数据"""
        data_list = []
        with open(file_path, 'r', encoding='utf-8', newline='') as f:
            reader = csv.DictReader(f)
            for row in reader:
                data_list.append(row)
        return data_list
    
    def _load_json_data(self, file_path: str) -> List[Dict[str, Any]]:
        """加载JSON数据"""
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        # 确保返回列表格式
        if isinstance(data, list):
            return data
        elif isinstance(data, dict):
            return [data]
        else:
            return []
    
    def _load_yaml_data(self, file_path: str) -> List[Dict[str, Any]]:
        """加载YAML数据"""
        with open(file_path, 'r', encoding='utf-8') as f:
            data = yaml.safe_load(f)
        # 确保返回列表格式
        if isinstance(data, list):
            return data
        elif isinstance(data, dict):
            return [data]
        else:
            return []
    
    def clear_cache(self) -> None:
        """清空数据缓存"""
        self.data_cache.clear()
